"""Find reusable public outputs for a goal. Never re-ingests an existing graph."""

from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

from cancer_output_atlas.first_look import first_look_hits, why_short
from cancer_output_atlas.ods import (
    ODS,
    ODS_ZH,
    TOOL_HINTS,
    has_posted_trial_results,
    ods_category,
    ods_zh,
    trial_results_status,
    trial_status_label,
)
from cancer_output_atlas.pipeline import records_from_json
from cancer_output_atlas.rank import parse_goal, rank_records
from cancer_output_atlas.schema import OutputRecord, classification_row

ABSTAIN_EN = "No public entries on this graph match that goal."
ABSTAIN_ZH = ABSTAIN_EN  # judges see English; *_zh kept for internal JSON


def _public_graph_ref(graph_path: str | None) -> str | None:
    """Omit internal filesystem paths (/app/..., /workspace/...) from public JSON."""
    if not graph_path:
        return None
    s = str(graph_path).replace("\\", "/")
    looks_fs = (
        s.startswith("/")
        or s.startswith("out/")
        or "/app/" in s
        or "/workspace/" in s
        or s.endswith(".json")
        or "link_graph" in s
    )
    if looks_fs:
        import os

        rev = (os.environ.get("K_REVISION") or "").strip()
        return rev or "baked"
    return s


SHANGHAI = ZoneInfo("Asia/Shanghai")

CALLABILITY = (
    "metadata_api",
    "runnable_workflow",
    "hosted_tool",
    "repo_only",
    "pointer_apply",
)

_CJK = re.compile(r"[\u3400-\u9fff]")
_WS = re.compile(r"\s+")
_CARD_SUMMARY_MAX = 220

CALLABILITY_ZH = {
    "metadata_api": "可查目录",
    "runnable_workflow": "可跑流程",
    "hosted_tool": "网页工具",
    "repo_only": "仓库",
    "pointer_apply": "需自己申请",
}

# Identifier schemes whose public *metadata* we already fetch (safety allow-list).
_METADATA_SCHEMES = frozenset(
    {
        "geo",
        "nct",
        "cbioportal",
        "gdc_project",
        "tcia_collection",
        "figshare",
        "europepmc",
    }
)
_WORKFLOW_SCHEMES = frozenset({"dockstore", "nfcore"})
_APPLY_HINTS = (
    "apply-yourself",
    "apply yourself",
    "apply for access",
    "controlled access",
    "user_applies",
    "dbgap",
    "controlled sequencing",
    "controlled file",
)


def _schemes(rec: OutputRecord) -> set[str]:
    return {i.scheme for i in rec.identifiers}


def _blob(rec: OutputRecord) -> str:
    return " ".join(
        p
        for p in (
            rec.title,
            rec.summary,
            rec.kind,
            rec.landing_url,
            rec.license_hint or "",
            " ".join(rec.identifier_values()),
        )
        if p
    ).lower()


def callability_of(rec: OutputRecord) -> str:
    """Tag how (if at all) this node can be turned into a result.

    Uses observed kind / identifiers / reuse flags and the existing
    one-click metadata-API allow-list. Does not invent IDs.
    """
    schemes = _schemes(rec)
    reuse = rec.reuse or {}
    landing = rec.landing_url or ""
    host = (urlparse(landing).hostname or "").lower()
    text = _blob(rec)

    if "gs_uri" in schemes or landing.startswith("gs://"):
        return "pointer_apply"
    if reuse.get("apply_yourself"):
        return "pointer_apply"
    if rec.source_status == "pointer" and not reuse.get("listing_site"):
        return "pointer_apply"
    if any(h in text for h in _APPLY_HINTS) and rec.source_status == "pointer":
        return "pointer_apply"

    if rec.kind == "workflow" or schemes & _WORKFLOW_SCHEMES:
        return "runnable_workflow"

    if rec.kind == "software" and "github" in schemes and (
        "github.com" in host or "github.com" in landing
    ):
        return "repo_only"

    if reuse.get("listing_site"):
        return "hosted_tool"

    if schemes & _METADATA_SCHEMES:
        return "metadata_api"

    if rec.kind == "software" or any(h in text for h in TOOL_HINTS):
        if "github.com" in host:
            return "repo_only"
        return "hosted_tool"

    if rec.source_status == "pointer":
        return "pointer_apply"
    return "metadata_api"


def _card_summary(rec: OutputRecord, *, limit: int = _CARD_SUMMARY_MAX) -> str:
    """English blurb from the record summary. Empty if missing, CJK-heavy, or a title echo."""
    raw = _WS.sub(" ", (rec.summary or "").strip())
    if not raw:
        return ""
    title = _WS.sub(" ", (rec.title or "").strip())
    if title and raw.casefold() == title.casefold():
        return ""
    cjk_n = len(_CJK.findall(raw))
    if cjk_n and (cjk_n >= 8 or cjk_n / max(len(raw), 1) >= 0.15):
        return ""
    if len(raw) <= limit:
        return raw
    cut = raw[:limit]
    for sep in (". ", "? ", "! ", "; "):
        idx = cut.rfind(sep)
        if idx >= limit // 2:
            return cut[: idx + 1].rstrip()
    idx = cut.rfind(" ")
    if idx >= int(limit * 0.6):
        return cut[:idx].rstrip(" ,;:") + "…"
    return cut.rstrip() + "…"


def _row_from_ranked(row: dict[str, Any], rec: OutputRecord) -> dict[str, Any]:
    if rec.callability in CALLABILITY:
        tag = rec.callability
    else:
        tag = callability_of(rec)
    ods = ods_category(
        kind=rec.kind,
        title=rec.title,
        summary=rec.summary,
        ids=[i.key() for i in rec.identifiers],
        landing_url=rec.landing_url,
        reuse=rec.reuse,
    )
    zh = ods_zh(ods)
    hit = {
        "output_id": row.get("output_id") or rec.output_id,
        "kind": row.get("kind") or rec.kind,
        "ods": ods,
        "ods_zh": zh,
        "title": row.get("title") or rec.title,
        "landing_url": row.get("landing_url") or rec.landing_url,
        "ids": list(row.get("ids") or [i.key() for i in rec.identifiers]),
        "why": row.get("why") or "",
        "goal_overlap": list(row.get("goal_overlap") or []),
        "profile_fit": row.get("profile_fit"),
        "rank_score": row.get("rank_score"),
        "callability": tag,
        "callability_zh": CALLABILITY_ZH[tag],
        "source_status": rec.source_status,
    }
    blurb = _card_summary(rec)
    if blurb:
        hit["summary"] = blurb
    hit["why_short"] = why_short(hit)
    schemes = {i.scheme for i in rec.identifiers}
    if rec.kind == "trial" or "nct" in schemes:
        posted = trial_results_status(title=rec.title, summary=rec.summary, reuse=rec.reuse)
        hit["has_results"] = posted
        hit["overall_status"] = (rec.reuse or {}).get("overall_status") or (rec.reuse or {}).get("overallStatus")
        hit["trial_status_label"] = trial_status_label(
            title=rec.title, summary=rec.summary, reuse=rec.reuse
        )
    return hit


def merge_catalog_pointers(records: list[OutputRecord]) -> list[OutputRecord]:
    """Add catalog_notes pointers (Arc VCC, apply-yourself cards) if missing.

    Public landing only. Never lists or downloads gs://. Safe when fixtures
    are absent (Cloud Run ships catalog_notes.json next to the app).
    """
    from cancer_output_atlas.classify import classify
    from cancer_output_atlas.sources.pointer import load_all_catalog_notes

    have: set[str] = set()
    for rec in records:
        have.add(rec.output_id)
        have.add(rec.landing_url or "")
        have.update(rec.all_id_keys())
        have.update(rec.identifier_values())
    extra: list[OutputRecord] = []
    for rec in load_all_catalog_notes():
        keys = {rec.output_id, rec.landing_url or "", *rec.all_id_keys(), *rec.identifier_values()}
        if keys & have:
            continue
        if rec.classification is None:
            classify(rec)
        extra.append(rec)
        have |= keys
    if not extra:
        return records
    return list(records) + extra



def _landing_key(url: str) -> str:
    return (url or "").strip().rstrip("/").lower()


def _scheme_specificity(rec: OutputRecord) -> int:
    schemes = {i.scheme for i in rec.identifiers}
    if schemes - {"url"}:
        return 1
    return 0


def _include_in_find(rec: OutputRecord, ods: str) -> bool:
    schemes = {i.scheme for i in rec.identifiers}
    if rec.kind == "trial" or "nct" in schemes:
        # Keep protocol-only NCT searchable; cards carry trial_status_label.
        return ods == "trial_result"
    if rec.kind == "publication" and ods != "method":
        return False
    return True


def _empty_find_payload(goal: str, *, top: int, graph_path: str | None) -> dict[str, Any]:
    now_utc = datetime.now(timezone.utc)
    now_sh = now_utc.astimezone(SHANGHAI)
    ts = now_utc.strftime("%Y-%m-%dT%H:%M:%SZ")
    try:
        from cancer_output_atlas.model_config import provider as model_provider
        provider_name = model_provider()
    except Exception:
        provider_name = "none"
    return {
        "goal": goal,
        "generated_at_utc": ts,
        "generated_at_shanghai": now_sh.strftime("%Y-%m-%dT%H:%M:%S+08:00"),
        "graph": _public_graph_ref(graph_path),
        "top_per_ods": top,
        "matched_total": 0,
        "matched_by_ods": {},
        "shown_total": 0,
        "by_ods": {},
        "first_look": [],
        "abstain": f"{ABSTAIN_EN} Snapshot {ts}.",
        "parser": "none",
        "model_provider": provider_name,
        "callability": {
            "software_tool_method": {k: 0 for k in CALLABILITY},
            "all_shown": {k: 0 for k in CALLABILITY},
        },
        "agent": "none",
        "strands_available": False,
        "strands_version": None,
        "tools_used": [],
        "agent_trace": [],
    }


def find_by_goal(
    records: list[OutputRecord],
    goal: str,
    *,
    top: int = 8,
    graph_path: str | None = None,
) -> dict[str, Any]:
    """Rank existing nodes against *goal* and group top hits by ODS."""
    goal = (goal or "").strip()
    if not goal:
        return _empty_find_payload("", top=top, graph_path=graph_path)
    records = merge_catalog_pointers(records)
    query = parse_goal(goal)
    ranked = rank_records(records, goal=goal, top=None, query_first=True, query=query)
    by_id = {r.output_id: r for r in records}
    grouped: dict[str, list[dict[str, Any]]] = {slug: [] for slug in ODS}
    seen_landings: dict[str, tuple[int, str]] = {}
    pending: list[tuple[dict[str, Any], OutputRecord]] = []
    for row in ranked:
        rec = by_id.get(row["output_id"])
        if rec is None:
            continue
        hit = _row_from_ranked(row, rec)
        if not _include_in_find(rec, hit["ods"]):
            continue
        pending.append((hit, rec))
    chosen: list[dict[str, Any]] = []
    for hit, rec in pending:
        key = _landing_key(hit.get("landing_url") or rec.landing_url or "")
        if not key:
            key = "id:" + (hit.get("output_id") or rec.output_id)
        spec = _scheme_specificity(rec)
        prev = seen_landings.get(key)
        if prev is None:
            seen_landings[key] = (len(chosen), spec)
            chosen.append(hit)
            continue
        idx, prev_spec = prev
        if spec > prev_spec:
            chosen[idx] = hit
            seen_landings[key] = (idx, spec)
    for hit in chosen:
        grouped[hit["ods"]].append(hit)

    matched_by_ods = {slug: len(grouped[slug]) for slug in ODS if grouped[slug]}
    matched_total = sum(matched_by_ods.values())

    by_ods: dict[str, list[dict[str, Any]]] = {}
    for slug in ODS:
        hits = grouped[slug]
        if not hits:
            continue
        if slug == "trial_result":
            hits.sort(
                key=lambda h: (
                    0 if h.get("has_results") is True else 1,
                    -(h.get("rank_score") or 0),
                )
            )
        for i, hit in enumerate(hits, start=1):
            hit["rank_in_ods"] = i
        by_ods[slug] = hits[:top]

    stm_counts: Counter[str] = Counter()
    all_counts: Counter[str] = Counter()
    for slug, hits in by_ods.items():
        for hit in hits:
            all_counts[hit["callability"]] += 1
            if slug in {"software", "tool", "method"}:
                stm_counts[hit["callability"]] += 1

    now_utc = datetime.now(timezone.utc)
    now_sh = now_utc.astimezone(SHANGHAI)
    first = first_look_hits(by_ods, 3)
    empty = not first and not by_ods
    try:
        from cancer_output_atlas.model_config import provider as model_provider
        provider_name = model_provider()
    except Exception:
        provider_name = "none"
    return {
        "goal": goal,
        "generated_at_utc": now_utc.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "generated_at_shanghai": now_sh.strftime("%Y-%m-%dT%H:%M:%S+08:00"),
        "graph": _public_graph_ref(graph_path),
        "top_per_ods": top,
        "matched_total": matched_total,
        "matched_by_ods": matched_by_ods,
        "shown_total": sum(len(v) for v in by_ods.values()),
        "by_ods": by_ods,
        "first_look": first,
        "abstain": (f"{ABSTAIN_EN} Snapshot {now_utc.strftime('%Y-%m-%dT%H:%M:%SZ')}.") if empty else None,
        "parser": query.source,
        "model_provider": provider_name,
        "callability": {
            "software_tool_method": {k: stm_counts.get(k, 0) for k in CALLABILITY},
            "all_shown": {k: all_counts.get(k, 0) for k in CALLABILITY},
        },
        "agent": "none",
        "strands_available": False,
        "strands_version": None,
        "tools_used": [],
        "agent_trace": [],
    }


def load_graph_records(graph_path: Path) -> list[OutputRecord]:
    payload = json.loads(graph_path.read_text(encoding="utf-8"))
    return records_from_json(payload)


def find_from_graph(graph_path: Path, goal: str, *, top: int = 8) -> dict[str, Any]:
    records = load_graph_records(graph_path)
    return find_by_goal(records, goal, top=top, graph_path=str(graph_path))


def format_find_text(payload: dict[str, Any]) -> str:
    lines = [
        f"goal\t{payload['goal']}",
        f"generated_at_utc\t{payload['generated_at_utc']}",
        f"generated_at_shanghai\t{payload['generated_at_shanghai']}",
        "",
    ]
    for slug in ODS:
        hits = payload["by_ods"].get(slug) or []
        lines.append(f"## {slug} / {ODS_ZH[slug]} ({len(hits)})")
        if not hits:
            lines.append("(none)")
            lines.append("")
            continue
        for hit in hits:
            ids = ",".join(hit.get("ids") or [])
            lines.append(
                f"{hit.get('rank_in_ods')}\t{hit['callability']}\t{hit['kind']}\t"
                f"{ids}\t{hit['title']}\t{hit['landing_url']}"
            )
            lines.append(f"  why: {hit['why']}")
        lines.append("")
    stm = payload["callability"]["software_tool_method"]
    lines.append("## callability (software/tool/method shown)")
    for k in CALLABILITY:
        lines.append(f"{k}\t{stm.get(k, 0)}")
    lines.append("")
    return "\n".join(lines)


def write_goal_find_json(payload: dict[str, Any], path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return path


def _xlsx_module():
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment
    except ImportError as exc:  # pragma: no cover - installed in demo env
        raise RuntimeError("openpyxl is required to write xlsx") from exc
    return Workbook, Font, Alignment


def write_goal_find_xlsx(payload: dict[str, Any], path: Path) -> Path:
    Workbook, Font, Alignment = _xlsx_module()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "goal"
    ws.append(["字段", "值"])
    ws.append(["goal", payload.get("goal") or ""])
    ws.append(["generated_at_utc", payload.get("generated_at_utc") or ""])
    ws.append(["generated_at_shanghai", payload.get("generated_at_shanghai") or ""])
    ws.append(["graph", payload.get("graph") or ""])
    ws.append(["top_per_ods", payload.get("top_per_ods")])
    ws["A1"].font = Font(bold=True)
    ws["B1"].font = Font(bold=True)
    fl = wb.create_sheet("xiankan")
    fl.title = "先看这几条"
    fl.append(["标题", "短理由", "落地页", "可调用性", "类型", "output_id"])
    for cell in fl[1]:
        cell.font = Font(bold=True)
    for hit in payload.get("first_look") or []:
        fl.append([
            hit.get("title"),
            hit.get("why_short") or hit.get("why"),
            hit.get("landing_url"),
            hit.get("callability_zh"),
            hit.get("ods_zh"),
            hit.get("output_id"),
        ])

    headers = [
        "ods",
        "ods_zh",
        "rank_in_ods",
        "kind",
        "title",
        "landing_url",
        "ids",
        "why",
        "callability",
        "callability_zh",
        "rank_score",
        "profile_fit",
        "output_id",
    ]
    zh_headers = [
        "ODS",
        "类型",
        "组内名次",
        "种类",
        "标题",
        "落地页",
        "标识符",
        "理由",
        "可调用性",
        "可调用性（中文）",
        "分数",
        "profile_fit",
        "output_id",
    ]
    for slug in ODS:
        ws_ods = wb.create_sheet(ODS_ZH[slug])
        ws_ods.append(zh_headers)
        for cell in ws_ods[1]:
            cell.font = Font(bold=True)
        for hit in payload["by_ods"].get(slug) or []:
            ws_ods.append(
                [
                    hit.get("ods"),
                    hit.get("ods_zh"),
                    hit.get("rank_in_ods"),
                    hit.get("kind"),
                    hit.get("title"),
                    hit.get("landing_url"),
                    "|".join(hit.get("ids") or []),
                    hit.get("why"),
                    hit.get("callability"),
                    hit.get("callability_zh"),
                    hit.get("rank_score"),
                    hit.get("profile_fit"),
                    hit.get("output_id"),
                ]
            )
            ws_ods.cell(ws_ods.max_row, 6).alignment = Alignment(wrap_text=True)
        ws_ods.column_dimensions["E"].width = 48
        ws_ods.column_dimensions["F"].width = 40
        ws_ods.column_dimensions["H"].width = 40

    ws_c = wb.create_sheet("callability")
    ws_c.append(["范围", "可调用性", "可调用性（中文）", "条数"])
    for cell in ws_c[1]:
        cell.font = Font(bold=True)
    stm = payload["callability"]["software_tool_method"]
    shown = payload["callability"]["all_shown"]
    for scope, counts in (
        ("software/tool/method", stm),
        ("all_shown", shown),
    ):
        for tag in CALLABILITY:
            ws_c.append([scope, tag, CALLABILITY_ZH[tag], counts.get(tag, 0)])

    wb.save(path)
    return path


def write_resources_xlsx(records: list[OutputRecord], path: Path) -> Path:
    """Full classification grouped by ODS (not goal-ranked)."""
    Workbook, Font, Alignment = _xlsx_module()
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws0 = wb.active
    ws0.title = "summary"
    counts: dict[str, int] = defaultdict(int)
    by_ods: dict[str, list[dict[str, Any]]] = {s: [] for s in ODS}
    for rec in records:
        if rec.source_status == "skipped":
            continue
        row = classification_row(rec)
        row["callability"] = callability_of(rec)
        row["callability_zh"] = CALLABILITY_ZH[row["callability"]]
        by_ods[row["ods"]].append(row)
        counts[row["ods"]] += 1
    ws0.append(["ods", "ods_zh", "count"])
    for cell in ws0[1]:
        cell.font = Font(bold=True)
    for slug in ODS:
        ws0.append([slug, ODS_ZH[slug], counts[slug]])

    headers = [
        "ods",
        "ods_zh",
        "kind",
        "title",
        "landing_url",
        "ids",
        "callability",
        "callability_zh",
        "profile_fit",
        "output_id",
    ]
    zh_headers = [
        "ODS",
        "类型",
        "种类",
        "标题",
        "落地页",
        "标识符",
        "可调用性",
        "可调用性（中文）",
        "profile_fit",
        "output_id",
    ]
    for slug in ODS:
        ws = wb.create_sheet(ODS_ZH[slug])
        ws.append(zh_headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in by_ods[slug]:
            ws.append(
                [
                    row.get("ods"),
                    row.get("ods_zh"),
                    row.get("kind"),
                    row.get("title"),
                    row.get("landing_url"),
                    "|".join(row.get("ids") or []),
                    row.get("callability"),
                    row.get("callability_zh"),
                    row.get("profile_fit"),
                    row.get("output_id"),
                ]
            )
        ws.column_dimensions["D"].width = 48
        ws.column_dimensions["E"].width = 40
    wb.save(path)
    return path


